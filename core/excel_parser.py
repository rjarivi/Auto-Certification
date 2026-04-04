import os
import openpyxl
from config import REQUIRED_COLUMNS


def parse_excel(filepath):
    """
    Parse an Excel file and return rows as list of dicts.
    Returns: {rows: [...], errors: [...], columns: [...]}
    """
    errors = []
    rows = []

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active

        raw_headers = [str(cell.value).strip() if cell.value else '' for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        headers = raw_headers

        # Case-insensitive column matching
        header_map = {}
        for required in REQUIRED_COLUMNS:
            for i, h in enumerate(headers):
                if h.lower() == required.lower():
                    header_map[required] = i
                    break

        missing = [col for col in REQUIRED_COLUMNS if col not in header_map]
        if missing:
            errors.append(f"Missing required columns: {', '.join(missing)}. Found: {', '.join(h for h in headers if h)}")
            wb.close()
            return {'rows': [], 'errors': errors, 'columns': headers}

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if all(cell is None for cell in row):
                continue
            record = {}
            for col_name, col_idx in header_map.items():
                val = row[col_idx] if col_idx < len(row) else None
                record[col_name] = str(val).strip() if val is not None else ''
            if not record.get('Email') or '@' not in record.get('Email', ''):
                errors.append(f"Row {row_idx}: Invalid or missing email '{record.get('Email', '')}'")
                continue
            if not record.get('Name'):
                errors.append(f"Row {row_idx}: Missing name, skipping.")
                continue
            rows.append(record)

        wb.close()

    except Exception as e:
        errors.append(f"Failed to read Excel file: {str(e)}")

    return {'rows': rows, 'errors': errors, 'columns': REQUIRED_COLUMNS}
